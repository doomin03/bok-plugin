"""Print a concise BOK figure-source workbook or sheet catalogue."""
import argparse
from pathlib import Path
import openpyxl
p=argparse.ArgumentParser()
p.add_argument("--workbook",required=True,type=Path); p.add_argument("--sheet"); p.add_argument("--preview-rows",type=int,default=12)
a=p.parse_args()
if not a.workbook.is_file(): p.error(f"Workbook does not exist: {a.workbook}")
b=openpyxl.load_workbook(a.workbook,read_only=True,data_only=False)
names=[a.sheet] if a.sheet else b.sheetnames
for n in names:
    if n not in b.sheetnames: p.error(f"Unknown sheet: {n}")
print(f"Workbook: {a.workbook.name} | sheets: {len(b.sheetnames)}")
for n in names:
    s=b[n]; print(f"\n[{n}] {s.max_row} rows x {s.max_column} columns"); count=0
    for i,row in enumerate(s.iter_rows(values_only=True),1):
        values=["" if v is None else " ".join(str(v).split()) for v in row]
        if not any(values): continue
        print(f"{i}: "+" | ".join(values)); count+=1
        if count>=a.preview_rows: break
    if not count: print("(no non-empty cells in preview)")
